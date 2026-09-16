from pathlib import Path

import pymupdf as fitz
from docx import Document
from fastapi import HTTPException, status


def extract_pdf_text(path: Path) -> str:
    try:
        document = fitz.open(path)

        try:
            pages = [
                page.get_text()
                for page in document
            ]
        finally:
            document.close()

        return "\n".join(pages).strip()

    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Unable to extract text from the PDF.",
        ) from exc


def extract_docx_text(path: Path) -> str:
    try:
        document = Document(path)

        content: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                content.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    content.append(" | ".join(cells))

        return "\n".join(content).strip()

    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Unable to extract text from the DOCX.",
        ) from exc


def extract_resume_text(path: Path) -> str:
    extension = path.suffix.lower()

    if extension == ".pdf":
        text = extract_pdf_text(path)

    elif extension == ".docx":
        text = extract_docx_text(path)

    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported resume format.",
        )

    if not text:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "No readable text was found in the resume. "
                "Scanned/image-only resumes are not supported yet."
            ),
        )

    return text